"""Extraction of financial metrics from CSV / XLSX / XLS / PDF files.

The strategy is table-oriented: every source is converted into rows of
(label, {period_column -> raw value}, source_ref). Labels are matched against
the metric dictionary; values are parsed with locale-aware number parsing.
No values are invented: metrics that are not found are reported as N/A.
"""
from __future__ import annotations

import io
from dataclasses import dataclass, field
from typing import Optional

from ..schemas import ExtractedValue, ExtractionResult, PeriodSelectionMeta
from . import extraction_headers as H
from . import extraction_ocr as EO
from . import extraction_pdf_text as EPT
from . import metrics as M
from . import ocr as OCR

# Resource caps: a hostile or pathological file (e.g. a 40k-row "sheet of
# junk") must not be fully materialized in memory or take minutes to parse.
# Real financial statements are tiny relative to these limits (findings 6, 26).
MAX_ROWS_PER_SHEET = 5000
MAX_COLS = 64
MAX_SHEETS = 20
MAX_PDF_PAGES = 60
TRUNCATION_WARNING = ("Файл усечён при разборе: обработаны не все строки/листы/страницы. "
                      "Финансовые отчёты обычно значительно меньше этих лимитов.")


class ScannedPdfError(Exception):
    """PDF has no extractable text. `ocr_warnings` (close wave, W4) carries
    OCR-path diagnostics for main.py's except branch to log."""
    def __init__(self, message: str, ocr_warnings: Optional[list[str]] = None):
        super().__init__(message)
        self.ocr_warnings = ocr_warnings or []


@dataclass
class _Row:
    label: str
    cells: list[str]
    source: str
    full_cells: list[str] = field(default_factory=list)  # F8: original column order


@dataclass
class _Table:
    header: list[str] = field(default_factory=list)
    rows: list[_Row] = field(default_factory=list)
    label_col_idx: int = 0
    code_col_idxs: frozenset[int] = frozenset()  # F2/S1: content-classified code columns


_HEADER_LABEL_TOKENS = ("наименование", "показател")
_LATEST_MARKERS = ("на конец", "end of")
_PREVIOUS_MARKERS = ("на начало", "beginning of")
_PERIOD_MARKER_PHRASES = _LATEST_MARKERS + _PREVIOUS_MARKERS + (
    "конец отчетного", "начало отчетного")
_KOD_TOKEN = "код"  # _label_based_column_roles' no-year mapping only — untouched by NEW-J
# T3 NEW-J: «код» alone missed «Строка 2025»; sync with extraction_headers.py.
_CODE_HEADER_TOKENS = ("код", "стр")


def _find_header_index(rows: list[list[str]]) -> Optional[int]:
    """Best header = the row (within the first 20 non-empty rows) whose
    non-first cells carry the most year-bearing cells. A title like
    «Отчёт ... за 2024 год» has its year in cell 0 and scores 0 (finding 3).

    Some statutory layouts (e.g. KZ RSBU) never put a year in the header at
    all — columns are labelled «На конец / начало отчётного периода»
    instead. When no row scores any year-bearing cells, fall back to
    recognizing the header by its label content: cell 0 reads as
    "наименование"/"показатель" and at least one other cell carries a
    period-marker phrase (finding F1)."""
    best_idx, best_count = None, 0
    for i, cells in enumerate(rows[:20]):
        count = sum(1 for c in cells[1:] if M.detect_periods([c]))
        if count > best_count:
            best_idx, best_count = i, count
    if best_count > 0:
        return best_idx
    for i, cells in enumerate(rows[:20]):
        if not cells:
            continue
        first = M.normalize_label(cells[0])
        if not any(tok in first for tok in _HEADER_LABEL_TOKENS):
            continue
        if any(any(marker in M.normalize_label(c) for marker in _PERIOD_MARKER_PHRASES)
               for c in cells[1:]):
            return i
    return best_idx


def _label_based_column_roles(data_header_cells: list[str]) -> dict[int, str]:
    """When a table's header has no year cells at all, map columns to
    latest/previous by header label content instead. «Код» columns are
    skipped — they hold line codes, not values. Columns carrying an
    explicit «на конец»/«на начало» (or «end of»/«beginning of») marker are
    assigned by that marker; any remaining, unmarked columns fall back to
    positional order (first → latest, second → previous).

    `data_header_cells` is already the label column's header cell excluded
    (see extraction_headers.data_columns) — callers pass the same slice
    used to build `_Row.cells`, so indices returned here line up with
    row.cells directly."""
    candidates: list[int] = []
    roles: dict[int, str] = {}
    for idx, cell in enumerate(data_header_cells):
        norm = M.normalize_label(cell)
        if _KOD_TOKEN in norm:
            continue
        candidates.append(idx)
        if any(m in norm for m in _LATEST_MARKERS):
            roles[idx] = "latest"
        elif any(m in norm for m in _PREVIOUS_MARKERS):
            roles[idx] = "previous"
    remaining_roles = [r for r in ("latest", "previous") if r not in roles.values()]
    for idx in candidates:
        if idx in roles or not remaining_roles:
            continue
        roles[idx] = remaining_roles.pop(0)
    return roles


def _rows_from_matrix(matrix: list[list], source_prefix: str) -> _Table:
    table = _Table()
    norm_rows: list[tuple[int, list[str]]] = []
    for i, row in enumerate(matrix):
        cells = ["" if c is None else str(c).strip() for c in row]
        if any(cells):
            norm_rows.append((i, cells))
    header_pos = _find_header_index([cells for _, cells in norm_rows])
    consumed: set[int] = set()
    if header_pos is not None:
        table.header, consumed = H.merge_header_block(norm_rows, header_pos)

    data_rows = [cells for pos, (_, cells) in enumerate(norm_rows)
                 if pos != header_pos and pos not in consumed]
    n_cols = max((len(cells) for _, cells in norm_rows), default=0)
    table.label_col_idx, table.code_col_idxs = H.classify_table(data_rows, table.header, n_cols)

    for pos, (i, cells) in enumerate(norm_rows):
        if pos == header_pos or pos in consumed:
            continue
        label = cells[table.label_col_idx] if table.label_col_idx < len(cells) else ""
        if label:
            table.rows.append(_Row(
                label=label, cells=H.data_columns(cells, table.label_col_idx),
                source=f"{source_prefix}, строка {i + 1}", full_cells=cells))
    return table


def _extract_from_tables(tables: list[_Table], full_text: str,
                         base_confidence_penalty: float = 0.0,
                         confidence_cap: Optional[float] = None) -> ExtractionResult:
    scale = M.detect_scale(full_text)
    currency = M.detect_currency(full_text)
    audited = M.detect_audited(full_text)

    # col_period (per table) is the single source of truth for "which
    # columns carry which period" used below for row value-assignment
    # (round 1, F3: «Код строки 2025» used to leak a phantom period here
    # despite the F2/S1 veto correctly excluding that column from data).
    #
    # period_objs (the overall latest/previous scan) is DELIBERATELY a
    # separate, wider scan of the table's WHOLE header — including the
    # label column's own header cell — not narrowed to header_data_cells
    # (round 2, NEW-2: round 1's restructure undeclaredly narrowed this,
    # silently dropping the period when a KZ statutory title embedded a
    # year in the label header, e.g. «Наименование показателя за 2024
    # год»; reverted to match pre-P7.T3/round-0 behavior). The F3 exclusion
    # (a code-header token, _CODE_HEADER_TOKENS — see NEW-J above) keeps
    # the phantom-period fix intact for anything code-like but not period-resolving.
    table_header_info: list[tuple[list[str], dict[int, str]]] = []
    period_objs: dict[str, M.Period] = {}
    for t in tables:
        # header_data_cells excludes the label column so indices line up
        # 1:1 with row.cells regardless of where the label column sits
        # (P7.T3a) — with label_col_idx == 0 this is exactly t.header[1:].
        header_data_cells = H.data_columns(t.header, t.label_col_idx) if t.header else []
        col_period: dict[int, str] = {}
        for idx, cell in enumerate(header_data_cells):
            # F2/S1, F3, NEW-J: never a period column despite a misleading
            # header year («Код строки 2025» / «Строка 2025») or none at all.
            normalized_cell = M.normalize_label(cell)
            if idx in t.code_col_idxs or any(tok in normalized_cell for tok in _CODE_HEADER_TOKENS):
                continue
            objs = M.detect_period_objects([cell])
            if objs:
                col_period[idx] = objs[0].label
        table_header_info.append((header_data_cells, col_period))

        for cell in t.header:
            if any(tok in M.normalize_label(cell) for tok in _CODE_HEADER_TOKENS):
                continue
            for p in M.detect_period_objects([cell]):
                period_objs.setdefault(p.label, p)
    ordered_periods = sorted(period_objs.values(), key=lambda p: p.sort_key, reverse=True)
    all_periods = [p.label for p in ordered_periods]
    latest_obj, previous_obj, period_meta = H.select_comparable_periods(ordered_periods)
    latest = latest_obj.label if latest_obj else None
    previous = previous_obj.label if previous_obj else None

    found: dict[str, ExtractedValue] = {}
    found_prev: dict[str, ExtractedValue] = {}
    warnings: list[str] = []

    # F7 (round 1): the chosen pair can be non-comparable (e.g. a quarter
    # vs. a full year, when no same-kind period exists — see
    # select_comparable_periods' fallback branch). That pair still feeds
    # the ratio engine and the confidence score's has-previous-period
    # bonus; period_selection.reason records *why* in meta, but a plain
    # ExtractionResult reader (the verify step) only sees `warnings`, so
    # the honesty law needs it here too.
    if latest_obj and previous_obj and latest_obj.kind != previous_obj.kind:
        warnings.append(
            f"Сравниваемые периоды разного типа — «{latest_obj.label}» ({latest_obj.kind}) и "
            f"«{previous_obj.label}» ({previous_obj.kind}): показатели могут быть несопоставимы."
        )

    for t, (header_data_cells, col_period) in zip(tables, table_header_info):
        has_year_columns = bool(col_period)

        # No year cells anywhere in the header: try to recognize it by label
        # content instead (e.g. KZ statutory «Наименование показателя / Код
        # строки / На конец / На начало отчётного периода», finding F1).
        col_role: dict[int, str] = {}
        label_header_recognized = False
        if not has_year_columns and t.header:
            label_cell = t.header[t.label_col_idx] if t.label_col_idx < len(t.header) else ""
            first_norm = M.normalize_label(label_cell)
            label_header_recognized = any(tok in first_norm for tok in _HEADER_LABEL_TOKENS)
            if label_header_recognized:
                col_role = _label_based_column_roles(header_data_cells)

        # Safety net: neither a year-bearing header nor a recognized label
        # header — we cannot reliably tell period columns from service
        # columns like «Код». Warn once and cap confidence for this table if
        # any matched row actually carries ≥2 numeric cells (finding F1b).
        safety_net_active = False
        if not has_year_columns and not label_header_recognized:
            for row in t.rows:
                if M.match_label(row.label) is None:
                    continue
                if sum(1 for c in row.cells if M.parse_number(c) is not None) >= 2:
                    safety_net_active = True
                    break
            if safety_net_active:
                note = ("Колонки периодов не распознаны — проверьте, что значения не взяты "
                        "из служебной колонки (например, «Код»).")
                if note not in warnings:
                    warnings.append(note)

        for row in t.rows:
            matched = M.match_label(row.label)
            if not matched:
                continue
            key, conf = matched
            conf = max(0.0, conf - base_confidence_penalty)
            if safety_net_active:
                conf = min(conf, 50.0)
            if confidence_cap is not None:
                conf = min(conf, confidence_cap)  # P7.T4: e.g. an OCR misread ("8" vs "3")
            values: dict[str, float | None] = {}
            for idx, raw in enumerate(row.cells):
                if idx in t.code_col_idxs:
                    continue  # F2/S1: content-classified code column — never a value,
                              # on every path (year-bearing, label-header, safety-net alike)
                num = M.parse_number(raw)
                if num is None:
                    continue
                if has_year_columns:
                    period = col_period.get(idx)
                    if period is None:
                        continue  # e.g. «Код» — a labeled non-year column is not data (finding 2)
                elif label_header_recognized:
                    role = col_role.get(idx)
                    if role is None:
                        continue  # «Код» column, or an extra column beyond latest/previous
                    period = role  # synthetic "latest"/"previous" key, resolved below
                else:
                    period = "latest"  # legacy fully-positional fallback (safety-net path)
                if period not in values:
                    values[period] = num
            # F8: built from the row's *original* column order (full_cells),
            # not label-then-data-columns — for a relocated label column
            # (P7.T3a) that order no longer matches the source document.
            # Identical to the pre-F8 "label | cells..." text whenever the
            # label sits at column 0 (every pre-P7.T3 fixture and the demo).
            snippet = " | ".join(c for c in row.full_cells if c)[:200]

            def put(store: dict, period_label: Optional[str], value: Optional[float]):
                if value is None:
                    return
                # Normalize sign BEFORE comparing to the stored value: the
                # stored value was already normalized when it was put(), so
                # comparing a raw (still-negative) incoming value against it
                # produces a false «Дублирующиеся значения» warning for the
                # same parenthesized figure appearing in two sections (F2).
                if key in M.EXPENSE_MAGNITUDE_METRICS and value < 0:
                    value = abs(value)
                    note = (f"Знак «{M.METRICS[key]['name']}» нормализован: значение в скобках "
                            "приведено к положительной величине расхода.")
                    if note not in warnings:
                        warnings.append(note)
                existing = store.get(key)
                if existing is not None and existing.value is not None:
                    if conf <= existing.confidence:
                        if existing.value != value:
                            warnings.append(
                                f"Дублирующиеся значения для «{M.METRICS[key]['name']}»: "
                                f"{existing.value} и {value}. Использовано более надёжное совпадение."
                            )
                        return
                    # new match is more confident (e.g. exact «Итого активы» later in the file)
                store[key] = ExtractedValue(
                    metric=key, original_label=row.label, value=value,
                    currency=currency, scale=scale, period=period_label,
                    source=row.source, confidence=conf, snippet=snippet,
                )

            latest_val = None
            if latest is not None and latest in values:
                latest_val = values[latest]
            elif "latest" in values:
                latest_val = values["latest"]
            if latest_val is not None:
                put(found, latest, latest_val)

            prev_val = None
            if previous is not None and previous in values:
                prev_val = values[previous]
            elif "previous" in values:
                prev_val = values["previous"]
            if prev_val is not None:
                put(found_prev, previous, prev_val)

    H.fill_missing_metrics(found, currency, scale, latest)
    warnings += H.near_total_miss_warning(found.values(), full_text)

    return ExtractionResult(
        upload_id="", periods=all_periods, latest_period=latest,
        previous_period=previous, currency=currency, scale=scale,
        audited=audited,
        values=list(found.values()),
        previous_values=list(found_prev.values()),
        warnings=warnings,
        period_selection=PeriodSelectionMeta(**period_meta),
    )


# ---------------------------------------------------------------------------
# Format-specific readers
# ---------------------------------------------------------------------------
def extract_from_csv(data: bytes) -> ExtractionResult:
    import csv
    text = data.decode("utf-8-sig", errors="replace")
    sample = text[:2048]
    delimiter = ";" if sample.count(";") > sample.count(",") else ","
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    matrix = [row for row in reader]
    table = _rows_from_matrix(matrix, "CSV")
    return _extract_from_tables([table], text)


def extract_from_xlsx(data: bytes) -> ExtractionResult:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    tables: list[_Table] = []
    full_text_parts: list[str] = []
    truncated = False
    for ws in wb.worksheets[:MAX_SHEETS]:
        matrix: list[list] = []
        for r_i, row in enumerate(ws.iter_rows(values_only=True)):
            if r_i >= MAX_ROWS_PER_SHEET:
                truncated = True
                break
            matrix.append(list(row[:MAX_COLS]))
        for row in matrix:
            full_text_parts.extend(str(c) for c in row if c is not None)
        tables.append(_rows_from_matrix(matrix, f"Лист «{ws.title}»"))
    if len(wb.worksheets) > MAX_SHEETS:
        truncated = True
    wb.close()
    result = _extract_from_tables(tables, " ".join(full_text_parts))
    if truncated:
        result.warnings.append(TRUNCATION_WARNING)
    return result


def extract_from_xls(data: bytes) -> ExtractionResult:
    import pandas as pd
    sheets = pd.read_excel(io.BytesIO(data), sheet_name=None, header=None,
                           dtype=str, nrows=MAX_ROWS_PER_SHEET)
    tables: list[_Table] = []
    full_text_parts: list[str] = []
    truncated = False
    sheet_names = list(sheets.keys())
    if len(sheet_names) > MAX_SHEETS:
        truncated = True
    for name in sheet_names[:MAX_SHEETS]:
        df = sheets[name]
        # pandas' nrows caps rows read but can't tell us whether the sheet
        # was actually longer than the cap — a sheet with exactly
        # MAX_ROWS_PER_SHEET rows read is treated as (possibly) truncated,
        # which is a conservative approximation (finding 6, 26).
        if len(df) >= MAX_ROWS_PER_SHEET:
            truncated = True
        matrix = df.iloc[:, :MAX_COLS].fillna("").values.tolist()
        for row in matrix:
            full_text_parts.extend(str(c) for c in row if c)
        tables.append(_rows_from_matrix(matrix, f"Лист «{name}»"))
    result = _extract_from_tables(tables, " ".join(full_text_parts))
    if truncated:
        result.warnings.append(TRUNCATION_WARNING)
    return result


def extract_from_pdf(data: bytes) -> ExtractionResult:
    import pdfplumber
    tables: list[_Table] = []
    text_parts: list[str] = []
    truncated = False
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        if len(pdf.pages) > MAX_PDF_PAGES:
            truncated = True
        for page_no, page in enumerate(pdf.pages[:MAX_PDF_PAGES], start=1):
            page_text = page.extract_text() or ""
            text_parts.append(page_text)
            page_tables = page.extract_tables() or []
            for t_no, raw_table in enumerate(page_tables, start=1):
                tables.append(_rows_from_matrix(
                    raw_table, f"PDF, стр. {page_no}, таблица {t_no}"))
            if not page_tables:
                # No ruling-line table — reconstruct by word position (see
                # extraction_pdf_text), else the older whitespace heuristic.
                matrix = EPT.words_to_matrix(page) or EO.lines_to_matrix(page_text)
                if matrix:
                    tables.append(_rows_from_matrix(matrix, f"PDF, стр. {page_no} (текст)"))
    full_text = "\n".join(text_parts)
    if not full_text.strip():  # no text layer (scan); byte-identical to pre-P7.T4 when OCR is off
        if not OCR.ocr_enabled():
            raise ScannedPdfError(
                "PDF не содержит текстового слоя (вероятно, это скан). "
                "Загрузите Excel/CSV либо PDF более высокого качества."
            )
        ocr_failed_message = ("OCR не смог распознать данные в этом PDF. "
                              "Загрузите Excel/CSV либо PDF более высокого качества.")
        ocr_texts, ocr_warnings = EO.ocr_pdf_pages(data)
        ocr_full_text = "\n".join(ocr_texts)
        if not ocr_full_text.strip():  # blank scan / poor quality / rasterization failure
            raise ScannedPdfError(ocr_failed_message, ocr_warnings=ocr_warnings)
        ocr_tables = [
            _rows_from_matrix(matrix, f"PDF, стр. {page_no} (распознано OCR)")
            for page_no, page_text in enumerate(ocr_texts, start=1)
            for matrix in [EO.lines_to_matrix(page_text)] if matrix
        ]
        result = _extract_from_tables(
            ocr_tables, ocr_full_text, base_confidence_penalty=15.0,
            confidence_cap=EO.CONFIDENCE_CAP)
        # Round-1: raw text can be non-empty GARBAGE (partial install -> eng-only).
        if not any(v.value is not None for v in result.values):
            raise ScannedPdfError(ocr_failed_message, ocr_warnings=ocr_warnings)
        result.warnings.append(EO.WARNING)
        result.warnings.extend(ocr_warnings)
        return result

    result = _extract_from_tables(tables, full_text, base_confidence_penalty=15.0)
    if truncated:
        result.warnings.append(TRUNCATION_WARNING)
    return result


def extract(data: bytes, kind: str) -> ExtractionResult:
    if kind == "csv":
        return extract_from_csv(data)
    if kind == "xlsx":
        return extract_from_xlsx(data)
    if kind == "xls":
        return extract_from_xls(data)
    if kind == "pdf":
        return extract_from_pdf(data)
    raise ValueError(f"Неподдерживаемый тип файла: {kind}")


def suggest_industry(result: ExtractionResult) -> Optional[str]:
    """Very rough heuristic industry hint based on available structure."""
    by_key = {v.metric: v for v in result.values}
    inv = by_key.get("inventory")
    capex = by_key.get("capital_expenditures")
    if inv and inv.value not in (None, 0) and capex and capex.value:
        return "manufacturing"
    if inv and inv.value not in (None, 0):
        return "retail"
    return None
